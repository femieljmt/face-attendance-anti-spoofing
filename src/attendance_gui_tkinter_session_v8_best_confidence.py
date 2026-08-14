
import csv
import os
import sys
import subprocess
import threading
from pathlib import Path
from datetime import datetime, timedelta
import tkinter as tk
from tkinter import ttk, messagebox

CLASSES = ["PCD", "TA", "Pempros", "PSD"]


class AttendanceGUI(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Face Attendance System - Liveness & Anti-Spoofing")
        self.geometry("1160x780")
        self.minsize(1080, 720)

        self.project_root = Path(__file__).resolve().parents[1]
        self.database_dir = self.project_root / "database"
        self.dataset_dir = self.project_root / "dataset" / "recognition"
        self.sessions_dir = self.database_dir / "attendance_sessions"
        self.backend_script = self.project_root / "src" / "final_attendance_lightweight_raspi.py"

        self.database_dir.mkdir(parents=True, exist_ok=True)
        self.dataset_dir.mkdir(parents=True, exist_ok=True)
        self.sessions_dir.mkdir(parents=True, exist_ok=True)

        self.student_csv = self.database_dir / "mahasiswa_gui.csv"
        self.attendance_log = self.database_dir / "absensi_runtime.csv"

        self.process = None
        self.reader_thread = None

        self.session_active = False
        self.waiting_backend_ready = False
        self.timer_started = False

        self.session_start = None
        self.backend_launch_time = None
        self.session_end = None
        self.session_id = None
        self.timer_job = None
        self.summary_file = None
        self.rejected_file = None
        self.session_raw_log = None
        self.session_class = None
        self.camera_test_process = None
        self.last_excel_file = None

        self._build_ui()
        self.repair_roster_classes()
        self.refresh_student_table()
        self.check_required_files(show_popup=False)
        self.update_system_indicators()

    # ============================================================
    # UI
    # ============================================================

    def _build_ui(self):
        header = ttk.Frame(self, padding=10)
        header.pack(fill="x")

        ttk.Label(header, text="Face Attendance System", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(
            header,
            text="Session-based attendance with liveness detection, anti-spoofing, face recognition, and automatic reporting.",
            font=("Segoe UI", 10)
        ).pack(anchor="w", pady=(2, 0))

        self.notebook = ttk.Notebook(self)
        self.notebook.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_registration = ttk.Frame(self.notebook, padding=10)
        self.tab_attendance = ttk.Frame(self.notebook, padding=10)
        self.tab_report = ttk.Frame(self.notebook, padding=10)
        self.tab_status = ttk.Frame(self.notebook, padding=10)

        self.notebook.add(self.tab_registration, text="Data Registration")
        self.notebook.add(self.tab_attendance, text="Attendance Recognition")
        self.notebook.add(self.tab_report, text="Session Report")
        self.notebook.add(self.tab_status, text="System Status")

        self._build_registration_tab()
        self._build_attendance_tab()
        self._build_report_tab()
        self._build_status_tab()

    def _build_registration_tab(self):
        form = ttk.LabelFrame(self.tab_registration, text="Student Registration", padding=12)
        form.pack(fill="x")

        ttk.Label(form, text="Kelas").grid(row=0, column=0, sticky="w", pady=4)
        self.reg_kelas = tk.StringVar(value="TA")
        ttk.Combobox(form, textvariable=self.reg_kelas, values=CLASSES, state="readonly", width=24).grid(row=0, column=1, sticky="w", pady=4)

        ttk.Label(form, text="NIM").grid(row=1, column=0, sticky="w", pady=4)
        self.reg_nim = tk.StringVar()
        ttk.Entry(form, textvariable=self.reg_nim, width=35).grid(row=1, column=1, sticky="w", pady=4)

        ttk.Label(form, text="Nama Mahasiswa").grid(row=2, column=0, sticky="w", pady=4)
        self.reg_name = tk.StringVar()
        ttk.Entry(form, textvariable=self.reg_name, width=45).grid(row=2, column=1, sticky="w", pady=4)

        self.create_folder_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(
            form,
            text="Buat folder dataset recognition untuk mahasiswa ini",
            variable=self.create_folder_var
        ).grid(row=3, column=1, sticky="w", pady=4)

        btns = ttk.Frame(form)
        btns.grid(row=4, column=1, sticky="w", pady=8)

        ttk.Button(btns, text="Simpan Data", command=self.save_student).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="Import Students from Dataset Folders", command=self.import_students_from_dataset).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="Refresh Tabel", command=self.refresh_student_table).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="Fix Kelas Roster", command=self.repair_roster_classes_clicked).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="Buka Folder Dataset", command=self.open_dataset_folder).pack(side="left")

        note = ttk.Label(
            self.tab_registration,
            text=(
                "Penting: folder dataset recognition tidak otomatis menjadi roster absensi. "
                "Klik 'Import Students from Dataset Folders' agar semua folder mahasiswa dimasukkan ke database/mahasiswa_gui.csv."
            ),
            foreground="#555555",
            wraplength=1000
        )
        note.pack(fill="x", pady=(10, 5))

        table_frame = ttk.LabelFrame(self.tab_registration, text="Registered Students / Roster", padding=8)
        table_frame.pack(fill="both", expand=True, pady=8)

        columns = ("kelas", "nim", "nama", "folder", "created_at")
        self.student_table = ttk.Treeview(table_frame, columns=columns, show="headings", height=14)

        headings = {
            "kelas": "Kelas",
            "nim": "NIM",
            "nama": "Nama",
            "folder": "Folder Dataset / Label",
            "created_at": "Created At",
        }

        widths = {
            "kelas": 90,
            "nim": 130,
            "nama": 260,
            "folder": 350,
            "created_at": 160,
        }

        for col in columns:
            self.student_table.heading(col, text=headings[col])
            self.student_table.column(col, width=widths[col], anchor="center" if col in ["kelas", "nim", "created_at"] else "w")

        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.student_table.yview)
        self.student_table.configure(yscrollcommand=yscroll.set)

        self.student_table.pack(side="left", fill="both", expand=True)
        yscroll.pack(side="right", fill="y")

    def _build_attendance_tab(self):
        config = ttk.LabelFrame(self.tab_attendance, text="Attendance Session Configuration", padding=12)
        config.pack(fill="x")

        self.run_kelas = tk.StringVar(value="TA")
        self.camera_index = tk.StringVar(value="0")
        self.width = tk.StringVar(value="640")
        self.height = tk.StringVar(value="480")
        self.fps = tk.StringVar(value="15")
        self.anti_threshold = tk.StringVar(value="0.03")
        self.rec_threshold = tk.StringVar(value="0.35")
        self.anti_frames = tk.StringVar(value="5")
        self.rec_frames = tk.StringVar(value="10")
        self.session_minutes = tk.StringVar(value="10")

        # Mode toggles
        self.liveness_active_var = tk.BooleanVar(value=True)
        self.antispoof_active_var = tk.BooleanVar(value=True)

        row = 0
        ttk.Label(config, text="Kelas").grid(row=row, column=0, sticky="w", pady=4)
        self.run_kelas_combo = ttk.Combobox(config, textvariable=self.run_kelas, values=CLASSES, state="readonly", width=18)
        self.run_kelas_combo.grid(row=row, column=1, sticky="w", pady=4)
        self.run_kelas_combo.bind("<<ComboboxSelected>>", self.on_run_class_changed)

        ttk.Label(config, text="Durasi Sesi (menit)").grid(row=row, column=2, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.session_minutes, width=8).grid(row=row, column=3, sticky="w", pady=4)

        ttk.Label(config, text="Camera").grid(row=row, column=4, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.camera_index, width=8).grid(row=row, column=5, sticky="w", pady=4)

        ttk.Label(config, text="Resolution").grid(row=row, column=6, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.width, width=7).grid(row=row, column=7, sticky="w", pady=4)
        ttk.Label(config, text="x").grid(row=row, column=8, sticky="w")
        ttk.Entry(config, textvariable=self.height, width=7).grid(row=row, column=9, sticky="w", pady=4)

        row += 1
        ttk.Label(config, text="FPS").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(config, textvariable=self.fps, width=8).grid(row=row, column=1, sticky="w", pady=4)

        ttk.Label(config, text="Anti-Spoof Threshold").grid(row=row, column=2, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.anti_threshold, width=8).grid(row=row, column=3, sticky="w", pady=4)

        ttk.Label(config, text="Recognition Threshold").grid(row=row, column=4, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.rec_threshold, width=8).grid(row=row, column=5, sticky="w", pady=4)

        row += 1
        ttk.Label(config, text="Anti-Spoof Frames").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(config, textvariable=self.anti_frames, width=8).grid(row=row, column=1, sticky="w", pady=4)

        ttk.Label(config, text="Recognition Frames").grid(row=row, column=2, sticky="w", padx=(20, 0), pady=4)
        ttk.Entry(config, textvariable=self.rec_frames, width=8).grid(row=row, column=3, sticky="w", pady=4)

        self.save_log_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(config, text="Simpan raw log session", variable=self.save_log_var).grid(row=row, column=4, columnspan=2, sticky="w", padx=(20, 0), pady=4)

        row += 1
        ttk.Label(config, text="Mode Pengujian").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Checkbutton(
            config,
            text="Liveness aktif",
            variable=self.liveness_active_var,
            command=self.update_system_indicators
        ).grid(row=row, column=1, sticky="w", pady=4)

        ttk.Checkbutton(
            config,
            text="Anti-spoofing aktif",
            variable=self.antispoof_active_var,
            command=self.update_system_indicators
        ).grid(row=row, column=2, columnspan=2, sticky="w", padx=(20, 0), pady=4)

        ttk.Label(
            config,
            text="Matikan keduanya untuk test recognition-only.",
            foreground="#555555"
        ).grid(row=row, column=4, columnspan=4, sticky="w", padx=(20, 0), pady=4)

        preset_frame = ttk.LabelFrame(self.tab_attendance, text="Mode Preset", padding=10)
        preset_frame.pack(fill="x", pady=(10, 0))

        ttk.Button(preset_frame, text="Mode Full System", command=self.set_mode_full_system).pack(side="left", padx=(0, 8))
        ttk.Button(preset_frame, text="Mode Recognition Only", command=self.set_mode_recognition_only).pack(side="left", padx=(0, 8))
        ttk.Button(preset_frame, text="Mode Liveness + Recognition", command=self.set_mode_liveness_recognition).pack(side="left", padx=(0, 8))
        ttk.Button(preset_frame, text="Mode Anti-Spoofing + Recognition", command=self.set_mode_antispoof_recognition).pack(side="left", padx=(0, 8))

        indicator_frame = ttk.LabelFrame(self.tab_attendance, text="System Indicators", padding=10)
        indicator_frame.pack(fill="x", pady=(10, 0))

        self.ind_camera = tk.StringVar(value="Camera: Not Tested")
        self.ind_model = tk.StringVar(value="Model: Checking")
        self.ind_database = tk.StringVar(value="Database: Checking")
        self.ind_session = tk.StringVar(value="Session: Stopped")
        self.ind_mode = tk.StringVar(value="Mode: Full System")

        ttk.Label(indicator_frame, textvariable=self.ind_camera, width=28).pack(side="left", padx=(0, 10))
        ttk.Label(indicator_frame, textvariable=self.ind_model, width=28).pack(side="left", padx=(0, 10))
        ttk.Label(indicator_frame, textvariable=self.ind_database, width=28).pack(side="left", padx=(0, 10))
        ttk.Label(indicator_frame, textvariable=self.ind_session, width=28).pack(side="left", padx=(0, 10))
        ttk.Label(indicator_frame, textvariable=self.ind_mode, width=36).pack(side="left", padx=(0, 10))

        session_status_frame = ttk.LabelFrame(self.tab_attendance, text="Session Status", padding=10)
        session_status_frame.pack(fill="x", pady=(10, 0))

        self.session_status = tk.StringVar(value="Session belum berjalan.")
        self.session_timer = tk.StringVar(value="00:00")
        ttk.Label(session_status_frame, textvariable=self.session_status, font=("Segoe UI", 10, "bold")).pack(side="left")
        ttk.Label(session_status_frame, text="Sisa waktu:", font=("Segoe UI", 10)).pack(side="left", padx=(30, 5))
        ttk.Label(session_status_frame, textvariable=self.session_timer, font=("Segoe UI", 12, "bold")).pack(side="left")

        buttons = ttk.Frame(self.tab_attendance)
        buttons.pack(fill="x", pady=10)

        ttk.Button(buttons, text="Test Camera", command=self.test_camera).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Start Session", command=self.start_session).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Start Timer Now", command=self.force_start_timer).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="End Session & Generate Report", command=self.end_session_manual).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Stop Backend Only", command=self.stop_attendance).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Open Session Folder", command=self.open_sessions_folder).pack(side="left", padx=(0, 8))
        ttk.Button(buttons, text="Check Required Files", command=lambda: self.check_required_files(show_popup=True)).pack(side="left")

        console_frame = ttk.LabelFrame(self.tab_attendance, text="Backend Console", padding=8)
        console_frame.pack(fill="both", expand=True)

        self.console = tk.Text(console_frame, height=18, wrap="word")
        self.console.pack(side="left", fill="both", expand=True)

        console_scroll = ttk.Scrollbar(console_frame, orient="vertical", command=self.console.yview)
        self.console.configure(yscrollcommand=console_scroll.set)
        console_scroll.pack(side="right", fill="y")

    def _build_report_tab(self):
        top = ttk.Frame(self.tab_report)
        top.pack(fill="x")

        ttk.Button(top, text="Generate Report From Current Log", command=self.generate_report_from_current_time_range).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Open Last Summary", command=self.open_last_summary).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Open Last Rejected", command=self.open_last_rejected).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Export Session Report to Excel", command=self.export_session_report_to_excel).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Open Session Folder", command=self.open_sessions_folder).pack(side="left", padx=(0, 8))
        ttk.Button(top, text="Open Raw Attendance Log", command=self.open_attendance_log).pack(side="left")

        note = ttk.Label(
            self.tab_report,
            text=(
                "Session summary dibuat tanpa duplikasi: satu mahasiswa hanya dicatat satu kali sebagai HADIR. "
                "Jika mahasiswa tidak ada di roster kelas, hasil valid dari backend dipindahkan ke rejected log."
            ),
            foreground="#555555",
            wraplength=980
        )
        note.pack(fill="x", pady=(8, 8))

        count_frame = ttk.LabelFrame(self.tab_report, text="Summary Count", padding=8)
        count_frame.pack(fill="x", pady=(0, 8))

        self.report_count_text = tk.StringVar(value="Total Mahasiswa: 0 | Hadir: 0 | Tidak Hadir: 0 | Ditolak: 0")
        ttk.Label(count_frame, textvariable=self.report_count_text, font=("Segoe UI", 10, "bold")).pack(anchor="w")

        summary_frame = ttk.LabelFrame(self.tab_report, text="Session Summary: Hadir / Tidak Hadir", padding=8)
        summary_frame.pack(fill="both", expand=True, pady=(0, 8))

        cols = ("status", "kelas", "nim", "nama", "waktu", "confidence", "spoof_prob", "keterangan")
        self.summary_table = ttk.Treeview(summary_frame, columns=cols, show="headings", height=12)

        headers = {
            "status": "Status",
            "kelas": "Kelas",
            "nim": "NIM",
            "nama": "Nama",
            "waktu": "Waktu",
            "confidence": "Conf.",
            "spoof_prob": "Spoof Prob.",
            "keterangan": "Keterangan",
        }

        widths = {
            "status": 110,
            "kelas": 80,
            "nim": 120,
            "nama": 220,
            "waktu": 150,
            "confidence": 80,
            "spoof_prob": 90,
            "keterangan": 300,
        }

        for col in cols:
            self.summary_table.heading(col, text=headers[col])
            self.summary_table.column(col, width=widths[col], anchor="center" if col not in ["nama", "keterangan"] else "w")

        self.summary_table.tag_configure("hadir", background="#d8f5d0")
        self.summary_table.tag_configure("tidak_hadir", background="#f8d7da")
        self.summary_table.tag_configure("netral", background="#ffffff")

        self.summary_table.pack(side="left", fill="both", expand=True)
        summary_scroll = ttk.Scrollbar(summary_frame, orient="vertical", command=self.summary_table.yview)
        self.summary_table.configure(yscrollcommand=summary_scroll.set)
        summary_scroll.pack(side="right", fill="y")

        rejected_frame = ttk.LabelFrame(self.tab_report, text="Rejected Attendance Attempts", padding=8)
        rejected_frame.pack(fill="both", expand=True)

        rcols = ("waktu", "kelas", "result", "status", "confidence", "spoof_prob", "reason")
        self.rejected_table = ttk.Treeview(rejected_frame, columns=rcols, show="headings", height=8)

        rheaders = {
            "waktu": "Waktu",
            "kelas": "Kelas",
            "result": "Result",
            "status": "Backend Status",
            "confidence": "Conf.",
            "spoof_prob": "Spoof Prob.",
            "reason": "Reason",
        }

        rwidths = {
            "waktu": 150,
            "kelas": 80,
            "result": 260,
            "status": 170,
            "confidence": 80,
            "spoof_prob": 90,
            "reason": 280,
        }

        for col in rcols:
            self.rejected_table.heading(col, text=rheaders[col])
            self.rejected_table.column(col, width=rwidths[col], anchor="center" if col not in ["result", "reason"] else "w")

        self.rejected_table.tag_configure("ditolak", background="#fff3cd")
        self.rejected_table.tag_configure("netral", background="#ffffff")

        self.rejected_table.pack(side="left", fill="both", expand=True)
        rejected_scroll = ttk.Scrollbar(rejected_frame, orient="vertical", command=self.rejected_table.yview)
        self.rejected_table.configure(yscrollcommand=rejected_scroll.set)
        rejected_scroll.pack(side="right", fill="y")

    def _build_status_tab(self):
        frame = ttk.LabelFrame(self.tab_status, text="Required Files", padding=10)
        frame.pack(fill="both", expand=True)

        self.status_text = tk.Text(frame, wrap="word")
        self.status_text.pack(fill="both", expand=True)

        btns = ttk.Frame(self.tab_status)
        btns.pack(fill="x", pady=8)

        ttk.Button(btns, text="Refresh Status", command=lambda: self.check_required_files(show_popup=False)).pack(side="left", padx=(0, 8))
        ttk.Button(btns, text="Open Project Folder", command=self.open_project_folder).pack(side="left")


    # ============================================================
    # MODE PRESET, INDICATORS, CAMERA TEST
    # ============================================================

    def get_mode_name(self):
        live = self.liveness_active_var.get()
        spoof = self.antispoof_active_var.get()

        if live and spoof:
            return "Full System"
        if (not live) and (not spoof):
            return "Recognition Only"
        if live and (not spoof):
            return "Liveness + Recognition"
        return "Anti-Spoofing + Recognition"

    def on_run_class_changed(self, event=None):
        kelas = self.run_kelas.get().strip()

        if kelas == "TA":
            # TA memakai:
            #   models/face_recognition_mobilenet_ta.h5
            #   models/class_indices_ta.json
            #
            # Recognition crop TA diatur pada backend final_attendance_lightweight_raspi.py.
            self.rec_threshold.set("0.25")
            self.rec_frames.set("10")
        else:
            self.rec_threshold.set("0.35")
            self.rec_frames.set("5")

        self.update_system_indicators()

    def set_mode_full_system(self):
        self.liveness_active_var.set(True)
        self.antispoof_active_var.set(True)
        self.update_system_indicators()

    def set_mode_recognition_only(self):
        self.liveness_active_var.set(False)
        self.antispoof_active_var.set(False)
        self.update_system_indicators()

    def set_mode_liveness_recognition(self):
        self.liveness_active_var.set(True)
        self.antispoof_active_var.set(False)
        self.update_system_indicators()

    def set_mode_antispoof_recognition(self):
        self.liveness_active_var.set(False)
        self.antispoof_active_var.set(True)
        self.update_system_indicators()

    def update_system_indicators(self):
        if not hasattr(self, "ind_model"):
            return

        missing = self.get_missing_files()
        if missing:
            self.ind_model.set("Model: Missing")
        else:
            self.ind_model.set("Model: Ready")

        roster_count = len(self.load_roster_for_class(self.run_kelas.get().strip()))
        if roster_count > 0:
            self.ind_database.set(f"Database: Ready ({roster_count})")
        else:
            self.ind_database.set("Database: Empty")

        if self.session_active:
            self.ind_session.set("Session: Running")
        elif self.waiting_backend_ready:
            self.ind_session.set("Session: Waiting Camera")
        else:
            self.ind_session.set("Session: Stopped")

        self.ind_mode.set(f"Mode: {self.get_mode_name()}")

    def test_camera(self):
        try:
            cam = int(self.camera_index.get())
            width = int(self.width.get())
            height = int(self.height.get())
            fps = int(self.fps.get())
        except ValueError:
            messagebox.showwarning("Input salah", "Camera, width, height, dan FPS harus berupa angka.")
            return

        if self.camera_test_process is not None and self.camera_test_process.poll() is None:
            messagebox.showinfo("Test Camera", "Jendela test camera masih berjalan. Tekan q di jendela kamera untuk keluar.")
            return

        probe_script = (
            "import cv2, sys;"
            f"cap=cv2.VideoCapture({cam});"
            f"cap.set(cv2.CAP_PROP_FRAME_WIDTH,{width});"
            f"cap.set(cv2.CAP_PROP_FRAME_HEIGHT,{height});"
            f"cap.set(cv2.CAP_PROP_FPS,{fps});"
            "ok=cap.isOpened();"
            "cap.release();"
            "sys.exit(0 if ok else 1)"
        )

        try:
            result = subprocess.run(
                [sys.executable, "-c", probe_script],
                cwd=str(self.project_root),
                timeout=6
            )
        except Exception as e:
            self.ind_camera.set("Camera: Not Ready")
            messagebox.showerror("Camera Error", f"Gagal mengecek kamera:\n{e}")
            return

        if result.returncode != 0:
            self.ind_camera.set("Camera: Not Ready")
            messagebox.showwarning("Camera Not Ready", f"Kamera index {cam} tidak bisa dibuka.")
            return

        self.ind_camera.set("Camera: Ready")

        viewer_script = f"""
import cv2
cap = cv2.VideoCapture({cam})
cap.set(cv2.CAP_PROP_FRAME_WIDTH, {width})
cap.set(cv2.CAP_PROP_FRAME_HEIGHT, {height})
cap.set(cv2.CAP_PROP_FPS, {fps})
if not cap.isOpened():
    print('Camera not opened')
    raise SystemExit(1)
print('Camera test running. Press q to close.')
while True:
    ret, frame = cap.read()
    if not ret:
        break
    cv2.putText(frame, 'TEST CAMERA - press q to close', (20, 35), cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0,255,255), 2)
    cv2.imshow('Test Camera', frame)
    if cv2.waitKey(1) & 0xFF == ord('q'):
        break
cap.release()
cv2.destroyAllWindows()
"""

        try:
            self.camera_test_process = subprocess.Popen(
                [sys.executable, "-u", "-c", viewer_script],
                cwd=str(self.project_root),
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )
        except Exception as e:
            messagebox.showerror("Camera Error", f"Kamera ready, tapi gagal membuka preview:\n{e}")

    # ============================================================
    # REGISTRATION
    # ============================================================

    def sanitize_folder_name(self, text):
        keep = []
        for ch in text:
            if ch.isalnum() or ch in ["_", "-"]:
                keep.append(ch)
            elif ch.isspace():
                keep.append("_")
        return "".join(keep).strip("_")

    def split_folder_to_nim_name(self, folder_name):
        if "_" in folder_name:
            parts = folder_name.split("_", 1)
            nim = parts[0].strip()
            name = parts[1].replace("_", " ").strip()
        else:
            parts = folder_name.split(" ", 1)
            nim = parts[0].strip()
            name = parts[1].strip() if len(parts) > 1 else folder_name

        return nim, name

    def save_student(self):
        kelas = self.reg_kelas.get().strip()
        nim = self.reg_nim.get().strip()
        name = self.reg_name.get().strip()

        if not kelas or not nim or not name:
            messagebox.showwarning("Data belum lengkap", "Kelas, NIM, dan nama wajib diisi.")
            return

        folder_name = f"{self.sanitize_folder_name(nim)}_{self.sanitize_folder_name(name)}"
        dataset_folder = self.dataset_dir / kelas / folder_name

        if self.create_folder_var.get():
            dataset_folder.mkdir(parents=True, exist_ok=True)

        if self.is_student_exists(kelas, nim):
            messagebox.showwarning("Data duplikat", f"NIM {nim} sudah terdaftar di kelas {kelas}.")
            return

        self.append_student(kelas, nim, name, dataset_folder)
        self.reg_nim.set("")
        self.reg_name.set("")
        self.refresh_student_table()
        self.update_system_indicators()

        messagebox.showinfo("Berhasil", f"Data mahasiswa disimpan.\nFolder: {dataset_folder}")

    def append_student(self, kelas, nim, name, dataset_folder):
        is_new_file = not self.student_csv.exists()

        with open(self.student_csv, "a", newline="", encoding="utf-8-sig") as f:
            fieldnames = ["kelas", "nim", "nama", "folder", "created_at"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)

            if is_new_file:
                writer.writeheader()

            writer.writerow({
                "kelas": kelas,
                "nim": nim,
                "nama": name,
                "folder": str(dataset_folder.relative_to(self.project_root)),
                "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            })

    def import_students_from_dataset(self):
        kelas = self.reg_kelas.get().strip()
        class_folder = self.dataset_dir / kelas

        if not class_folder.exists():
            messagebox.showwarning("Folder tidak ditemukan", f"Folder tidak ditemukan:\n{class_folder}")
            return

        imported = 0
        skipped = 0

        for item in sorted(class_folder.iterdir()):
            if not item.is_dir():
                continue

            nim, name = self.split_folder_to_nim_name(item.name)

            if not nim:
                skipped += 1
                continue

            if self.is_student_exists(kelas, nim):
                skipped += 1
                continue

            self.append_student(kelas, nim, name, item)
            imported += 1

        self.refresh_student_table()
        self.update_system_indicators()

        messagebox.showinfo(
            "Import selesai",
            f"Kelas: {kelas}\nImported: {imported}\nSkipped/duplikat: {skipped}"
        )

    def is_student_exists(self, kelas, nim):
        for row in self.load_roster_all():
            if row.get("kelas") == kelas and row.get("nim") == nim:
                return True
        return False

    def refresh_student_table(self):
        for item in self.student_table.get_children():
            self.student_table.delete(item)

        for row in self.load_roster_all():
            self.student_table.insert(
                "",
                "end",
                values=(
                    row.get("kelas", ""),
                    row.get("nim", ""),
                    row.get("nama", ""),
                    row.get("folder", ""),
                    row.get("created_at", "")
                )
            )

    def infer_class_from_folder(self, folder_text):
        folder_text = str(folder_text or "").replace("\\", "/").strip()
        parts = [p for p in folder_text.split("/") if p]

        # Pola utama:
        # dataset/recognition/TA/Nama_Mahasiswa
        # dataset/recognition/PCD/Nama_Mahasiswa
        for i, part in enumerate(parts):
            if part.lower() == "recognition" and i + 1 < len(parts):
                candidate = parts[i + 1].strip()
                for kelas in CLASSES:
                    if candidate.lower() == kelas.lower():
                        return kelas

        # Fallback: jika ada segmen folder yang sama dengan nama kelas.
        for part in parts:
            for kelas in CLASSES:
                if part.lower() == kelas.lower():
                    return kelas

        return ""

    def normalize_roster_row(self, row):
        normalized = dict(row)

        # Pastikan semua kolom dasar ada.
        for key in ["kelas", "nim", "nama", "folder", "created_at"]:
            normalized.setdefault(key, "")

        kelas = str(normalized.get("kelas", "") or "").strip()

        if kelas not in CLASSES:
            inferred = self.infer_class_from_folder(normalized.get("folder", ""))
            if inferred:
                normalized["kelas"] = inferred

        return normalized

    def write_roster_rows(self, rows, fieldnames=None):
        self.student_csv.parent.mkdir(parents=True, exist_ok=True)

        base_fields = ["kelas", "nim", "nama", "folder", "created_at"]

        if fieldnames is None:
            fieldnames = base_fields[:]

        for field in base_fields:
            if field not in fieldnames:
                fieldnames.append(field)

        # Tulis ulang CSV dengan kelas yang sudah diperbaiki.
        with open(self.student_csv, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                for field in fieldnames:
                    row.setdefault(field, "")
                writer.writerow(row)

    def repair_roster_classes(self):
        if not self.student_csv.exists():
            return 0, 0

        try:
            with open(self.student_csv, "r", newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                fieldnames = reader.fieldnames or ["kelas", "nim", "nama", "folder", "created_at"]
                rows_raw = list(reader)

            fixed_rows = []
            fixed_count = 0

            for row in rows_raw:
                old_kelas = str(row.get("kelas", "") or "").strip()
                fixed = self.normalize_roster_row(row)
                new_kelas = str(fixed.get("kelas", "") or "").strip()

                if old_kelas != new_kelas and new_kelas:
                    fixed_count += 1

                fixed_rows.append(fixed)

            if fixed_count > 0:
                self.write_roster_rows(fixed_rows, fieldnames=fieldnames)

            return fixed_count, len(fixed_rows)

        except Exception as e:
            messagebox.showerror("Error", f"Gagal memperbaiki kolom kelas roster:\n{e}")
            return 0, 0

    def repair_roster_classes_clicked(self):
        fixed_count, total_rows = self.repair_roster_classes()
        self.refresh_student_table()
        self.update_system_indicators()

        messagebox.showinfo(
            "Fix Kelas Roster",
            f"Perbaikan selesai.\n\n"
            f"Baris diperbaiki: {fixed_count}\n"
            f"Total baris roster: {total_rows}\n\n"
            f"Jika tabel belum berubah, tutup GUI lalu buka ulang."
        )

    def load_roster_all(self):
        rows = []

        if not self.student_csv.exists():
            return rows

        try:
            with open(self.student_csv, "r", newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                fieldnames = reader.fieldnames or ["kelas", "nim", "nama", "folder", "created_at"]
                changed = False

                for row in reader:
                    old_kelas = str(row.get("kelas", "") or "").strip()
                    fixed = self.normalize_roster_row(row)
                    new_kelas = str(fixed.get("kelas", "") or "").strip()

                    if old_kelas != new_kelas and new_kelas:
                        changed = True

                    rows.append(fixed)

            # Auto-fix permanen: jika kelas bisa dibaca dari folder,
            # langsung tulis ulang mahasiswa_gui.csv.
            if changed:
                self.write_roster_rows(rows, fieldnames=fieldnames)

        except Exception as e:
            messagebox.showerror("Error", f"Gagal membaca mahasiswa_gui.csv:\n{e}")

        return rows

    def load_roster_for_class(self, kelas):
        kelas = str(kelas or "").strip()
        return [
            r for r in self.load_roster_all()
            if str(r.get("kelas", "") or "").strip() == kelas
        ]

    # ============================================================
    # SESSION CONTROL
    # ============================================================

    def start_session(self):
        if self.session_active or self.waiting_backend_ready:
            messagebox.showwarning("Session aktif", "Session absensi masih berjalan atau menunggu backend siap.")
            return

        try:
            duration_minutes = float(self.session_minutes.get())
            if duration_minutes <= 0:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Durasi salah", "Durasi sesi harus berupa angka menit lebih dari 0.")
            return

        # Kunci kelas saat session dimulai.
        # Ini mencegah report memakai kelas lama/default jika combobox berubah atau tab berpindah.
        self.session_class = self.run_kelas.get().strip()

        roster = self.load_roster_for_class(self.session_class)

        if not roster:
            ok = messagebox.askyesno(
                "Roster kosong",
                "Belum ada mahasiswa pada kelas ini di database/mahasiswa_gui.csv.\n"
                "Klik Import Students from Dataset Folders di tab Data Registration.\n\n"
                "Tetap jalankan session?"
            )
            if not ok:
                return

        self.waiting_backend_ready = True
        self.timer_started = False
        self.backend_launch_time = datetime.now()
        self.session_id = self.backend_launch_time.strftime("%Y%m%d_%H%M%S") + "_" + self.session_class
        self.session_raw_log = self.sessions_dir / f"session_raw_{self.session_id}.csv"

        self.session_status.set(f"Menunggu backend dan kamera siap: {self.session_class} | Timer belum berjalan")
        self.session_timer.set("WAIT")
        self.update_system_indicators()

        self.console_delete()
        self.console_write(f"[SESSION] Kelas dikunci untuk session ini: {self.session_class}\n")
        self.console_write(f"[SESSION] Raw log session: {self.session_raw_log}\n\n")

        self.start_backend()

    def start_timer_when_ready(self):
        if self.timer_started:
            return

        try:
            duration_minutes = float(self.session_minutes.get())
        except ValueError:
            duration_minutes = 5.0

        self.session_active = True
        self.waiting_backend_ready = False
        self.timer_started = True

        self.session_start = datetime.now()
        self.session_end = self.session_start + timedelta(minutes=duration_minutes)

        self.session_status.set(
            f"Session berjalan: {self.session_class} | Mulai {self.session_start.strftime('%H:%M:%S')}"
        )
        self.session_timer.set(self.format_remaining_time())
        self.update_timer()
        self.update_system_indicators()

        self.console_write("\n[TIMER] Timer session dimulai setelah backend siap.\n")

    def force_start_timer(self):
        if self.timer_started:
            messagebox.showinfo("Timer sudah berjalan", "Timer session sudah berjalan.")
            return

        if self.process is None or self.process.poll() is not None:
            messagebox.showwarning("Backend belum berjalan", "Backend belum berjalan. Klik Start Session dulu.")
            return

        self.start_timer_when_ready()

    def end_session_manual(self):
        if not self.session_active and not self.waiting_backend_ready:
            messagebox.showinfo("Session belum aktif", "Tidak ada session yang sedang berjalan.")
            return

        self.finish_session(reason="manual")

    def finish_session(self, reason="timer"):
        if not self.session_active and not self.waiting_backend_ready:
            return

        self.session_active = False
        self.waiting_backend_ready = False
        self.timer_started = False

        if self.timer_job is not None:
            try:
                self.after_cancel(self.timer_job)
            except Exception:
                pass
            self.timer_job = None

        self.stop_attendance()

        actual_end = datetime.now()

        if self.session_start is None:
            self.session_start = self.backend_launch_time or actual_end

        self.session_status.set(f"Session selesai ({reason}) | {self.session_class} | {actual_end.strftime('%H:%M:%S')}")
        self.session_timer.set("00:00")
        self.update_system_indicators()

        summary_rows, rejected_rows = self.generate_session_report(
            kelas=self.session_class,
            start_time=self.session_start,
            end_time=actual_end,
            session_id=self.session_id
        )

        self.populate_report_tables(summary_rows, rejected_rows)

        # Auto-refresh: pindah otomatis ke tab Session Report setelah session selesai.
        if hasattr(self, "notebook"):
            self.notebook.select(self.tab_report)

        messagebox.showinfo(
            "Session selesai",
            "Rekap absensi selesai dibuat.\n\n"
            f"Summary:\n{self.summary_file}\n\n"
            f"Rejected:\n{self.rejected_file}"
        )

    def update_timer(self):
        if not self.session_active:
            return

        remaining = (self.session_end - datetime.now()).total_seconds()

        if remaining <= 0:
            self.finish_session(reason="timer")
            return

        self.session_timer.set(self.format_remaining_time())
        self.timer_job = self.after(1000, self.update_timer)

    def format_remaining_time(self):
        if self.session_end is None:
            return "00:00"

        remaining = max(0, int((self.session_end - datetime.now()).total_seconds()))
        minutes = remaining // 60
        seconds = remaining % 60
        return f"{minutes:02d}:{seconds:02d}"

    # ============================================================
    # BACKEND
    # ============================================================

    def build_backend_command(self):
        kelas_for_backend = self.session_class or self.run_kelas.get().strip()

        cmd = [
            sys.executable,
            "-u",
            str(self.backend_script),
            "--kelas", kelas_for_backend,
            "--camera", self.camera_index.get(),
            "--width", self.width.get(),
            "--height", self.height.get(),
            "--fps", self.fps.get(),
            "--anti_spoof_threshold", self.anti_threshold.get(),
            "--anti_spoof_frames", self.anti_frames.get(),
            "--recognition_frames", self.rec_frames.get(),
            "--recognition_threshold", self.rec_threshold.get(),
        ]

        if not self.liveness_active_var.get():
            cmd.append("--disable_liveness")

        if not self.antispoof_active_var.get():
            cmd.append("--disable_antispoof")

        if (not self.liveness_active_var.get()) and (not self.antispoof_active_var.get()):
            cmd.append("--recognition_only")

        if self.save_log_var.get():
            cmd.append("--save_log")
            cmd.extend(["--log_path", str(self.session_raw_log)])

        return cmd

    def start_backend(self):
        if self.process is not None and self.process.poll() is None:
            messagebox.showwarning("Masih berjalan", "Backend absensi masih berjalan.")
            return

        if not self.backend_script.exists():
            messagebox.showerror("File tidak ditemukan", f"Backend tidak ditemukan:\n{self.backend_script}")
            return

        missing = self.get_missing_files()
        if missing:
            messagebox.showwarning(
                "Ada file belum lengkap",
                "Beberapa file belum ditemukan:\n\n" + "\n".join(missing)
            )

        cmd = self.build_backend_command()

        # Jangan hapus console di sini karena start_session sudah menulis informasi kelas session.
        self.console_write("Menjalankan backend:\n")
        self.console_write(" ".join(cmd) + "\n\n")

        try:
            self.process = subprocess.Popen(
                cmd,
                cwd=str(self.project_root),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                bufsize=1
            )
        except Exception as e:
            messagebox.showerror("Gagal menjalankan backend", str(e))
            return

        self.reader_thread = threading.Thread(target=self.read_backend_output, daemon=True)
        self.reader_thread.start()

    def read_backend_output(self):
        if self.process is None or self.process.stdout is None:
            return

        for line in self.process.stdout:
            self.after(0, self.handle_backend_line, line)

        self.after(0, self.console_write, "\nBackend selesai.\n")

    def handle_backend_line(self, line):
        self.console_write(line)

        text = line.strip().lower()

        ready_signals = [
            "program berjalan",
            "q = keluar",
            "haar cascade aktif",
            "haar aktif",
            "state:",
            "liveness:",
            "real (",
            "spoofprob",
        ]

        if self.waiting_backend_ready and not self.timer_started:
            if any(signal in text for signal in ready_signals):
                self.start_timer_when_ready()

    def stop_attendance(self):
        if self.process is None or self.process.poll() is not None:
            self.console_write("Tidak ada backend yang sedang berjalan.\n")
            return

        self.process.terminate()
        self.console_write("\nMenghentikan backend...\n")

        try:
            self.process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            self.process.kill()
            self.console_write("Backend dipaksa berhenti.\n")

        self.update_system_indicators()

    # ============================================================
    # SESSION REPORT
    # ============================================================

    def parse_datetime(self, text):
        try:
            return datetime.strptime(text, "%Y-%m-%d %H:%M:%S")
        except Exception:
            return None

    def get_current_raw_log(self):
        if self.session_raw_log and self.session_raw_log.exists():
            return self.session_raw_log
        return self.attendance_log

    def read_raw_log_in_time_range(self, start_time, end_time, kelas):
        rows = []
        raw_log = self.get_current_raw_log()

        if not raw_log.exists():
            return rows

        try:
            with open(raw_log, "r", newline="", encoding="utf-8") as f:
                reader = csv.DictReader(f)

                for row in reader:
                    ts = self.parse_datetime(row.get("timestamp", ""))

                    if ts is None:
                        continue

                    if start_time <= ts <= end_time and row.get("kelas") == kelas:
                        rows.append(row)
        except Exception as e:
            messagebox.showerror("Error", f"Gagal membaca raw attendance log:\n{e}")

        return rows

    def possible_labels_for_student(self, student):
        labels = set()

        nim = student.get("nim", "").strip()
        nama = student.get("nama", "").strip()
        folder = student.get("folder", "").strip()

        if folder:
            labels.add(Path(folder).name)

        if nim:
            labels.add(nim)

        if nim and nama:
            labels.add(f"{self.sanitize_folder_name(nim)}_{self.sanitize_folder_name(nama)}")

        return labels

    def find_student_for_result(self, roster, result):
        result = (result or "").strip()

        if not result:
            return None

        for student in roster:
            nim = student.get("nim", "").strip()
            labels = self.possible_labels_for_student(student)

            if result in labels:
                return student

            if nim and result.startswith(nim):
                return student

        return None

    def safe_float(self, value, default=-1.0):
        try:
            if value is None or value == "":
                return default
            return float(value)
        except Exception:
            return default

    def generate_session_report(self, kelas, start_time, end_time, session_id):
        roster = self.load_roster_for_class(kelas)
        raw_rows = self.read_raw_log_in_time_range(start_time, end_time, kelas)

        present_by_nim = {}
        rejected_rows = []
        rejected_seen = set()

        for row in raw_rows:
            status = row.get("status", "").strip()
            result = row.get("result", "").strip()

            if status == "ABSENSI_VALID":
                student = self.find_student_for_result(roster, result)

                if student is not None:
                    nim = student.get("nim", "").strip()
                    current_conf = self.safe_float(row.get("recognition_confidence", ""), default=-1.0)

                    # Jika mahasiswa sudah pernah tercatat HADIR pada sesi yang sama,
                    # jangan buat baris duplikat. Namun simpan data dengan confidence tertinggi.
                    # Jadi rekap akhir tetap satu mahasiswa satu baris, tetapi nilai Conf. adalah yang terbaik.
                    existing = present_by_nim.get(nim)
                    existing_conf = self.safe_float(existing.get("confidence", ""), default=-1.0) if existing else -1.0

                    if existing is None or current_conf > existing_conf:
                        present_by_nim[nim] = {
                            "status": "HADIR",
                            "kelas": kelas,
                            "nim": nim,
                            "nama": student.get("nama", ""),
                            "waktu": row.get("timestamp", ""),
                            "confidence": row.get("recognition_confidence", ""),
                            "spoof_prob": row.get("spoof_probability", ""),
                            "keterangan": "Absensi valid. Duplikasi diabaikan, dipilih confidence tertinggi."
                        }
                else:
                    key = (row.get("timestamp", ""), result, status)
                    if key not in rejected_seen:
                        rejected_seen.add(key)
                        rejected_rows.append({
                            "waktu": row.get("timestamp", ""),
                            "kelas": kelas,
                            "result": result,
                            "status": status,
                            "confidence": row.get("recognition_confidence", ""),
                            "spoof_prob": row.get("spoof_probability", ""),
                            "reason": "DITOLAK_TIDAK_TERDAFTAR_DI_KELAS"
                        })

            else:
                key = (row.get("timestamp", ""), result, status)
                if key not in rejected_seen:
                    rejected_seen.add(key)
                    rejected_rows.append({
                        "waktu": row.get("timestamp", ""),
                        "kelas": kelas,
                        "result": result,
                        "status": status,
                        "confidence": row.get("recognition_confidence", ""),
                        "spoof_prob": row.get("spoof_probability", ""),
                        "reason": self.map_rejection_reason(status)
                    })

        summary_rows = []

        for student in roster:
            nim = student.get("nim", "").strip()

            if nim in present_by_nim:
                summary_rows.append(present_by_nim[nim])
            else:
                summary_rows.append({
                    "status": "TIDAK HADIR",
                    "kelas": kelas,
                    "nim": nim,
                    "nama": student.get("nama", ""),
                    "waktu": "",
                    "confidence": "",
                    "spoof_prob": "",
                    "keterangan": "Tidak tercatat hadir selama sesi."
                })

        self.write_session_files(session_id, summary_rows, rejected_rows)
        return summary_rows, rejected_rows

    def map_rejection_reason(self, status):
        mapping = {
            "DITOLAK_LIVENESS": "Liveness gagal",
            "DITOLAK_SPOOF": "Terdeteksi spoof / replay",
            "DITOLAK_TIDAK_DIKENAL": "Wajah tidak dikenal",
        }
        return mapping.get(status, "Ditolak oleh backend")

    def write_session_files(self, session_id, summary_rows, rejected_rows):
        self.sessions_dir.mkdir(parents=True, exist_ok=True)

        self.summary_file = self.sessions_dir / f"session_summary_{session_id}.csv"
        self.rejected_file = self.sessions_dir / f"session_rejected_{session_id}.csv"

        with open(self.summary_file, "w", newline="", encoding="utf-8") as f:
            fieldnames = ["status", "kelas", "nim", "nama", "waktu", "confidence", "spoof_prob", "keterangan"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for row in summary_rows:
                writer.writerow(row)

        with open(self.rejected_file, "w", newline="", encoding="utf-8") as f:
            fieldnames = ["waktu", "kelas", "result", "status", "confidence", "spoof_prob", "reason"]
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for row in rejected_rows:
                writer.writerow(row)

    def populate_report_tables(self, summary_rows, rejected_rows):
        for item in self.summary_table.get_children():
            self.summary_table.delete(item)

        hadir_count = 0
        tidak_hadir_count = 0

        for row in summary_rows:
            status = row.get("status", "")
            if status == "HADIR":
                hadir_count += 1
                tag = "hadir"
            elif status == "TIDAK HADIR":
                tidak_hadir_count += 1
                tag = "tidak_hadir"
            else:
                tag = "netral"

            self.summary_table.insert(
                "",
                "end",
                values=(
                    row.get("status", ""),
                    row.get("kelas", ""),
                    row.get("nim", ""),
                    row.get("nama", ""),
                    row.get("waktu", ""),
                    row.get("confidence", ""),
                    row.get("spoof_prob", ""),
                    row.get("keterangan", "")
                ),
                tags=(tag,)
            )

        for item in self.rejected_table.get_children():
            self.rejected_table.delete(item)

        for row in rejected_rows:
            self.rejected_table.insert(
                "",
                "end",
                values=(
                    row.get("waktu", ""),
                    row.get("kelas", ""),
                    row.get("result", ""),
                    row.get("status", ""),
                    row.get("confidence", ""),
                    row.get("spoof_prob", ""),
                    row.get("reason", "")
                ),
                tags=("ditolak",)
            )

        if hasattr(self, "report_count_text"):
            self.report_count_text.set(
                f"Total Mahasiswa: {len(summary_rows)} | "
                f"Hadir: {hadir_count} | "
                f"Tidak Hadir: {tidak_hadir_count} | "
                f"Ditolak: {len(rejected_rows)}"
            )

    def generate_report_from_current_time_range(self):
        if self.session_start is None:
            messagebox.showinfo(
                "Belum ada session",
                "Belum ada waktu session aktif. Jalankan Start Session terlebih dahulu."
            )
            return

        end_time = datetime.now()
        report_class = self.session_class or self.run_kelas.get().strip()
        session_id = self.session_id or (self.session_start.strftime("%Y%m%d_%H%M%S") + "_" + report_class)

        summary_rows, rejected_rows = self.generate_session_report(
            kelas=report_class,
            start_time=self.session_start,
            end_time=end_time,
            session_id=session_id
        )

        self.populate_report_tables(summary_rows, rejected_rows)

        # Auto-refresh: pindah otomatis ke tab Session Report setelah session selesai.
        if hasattr(self, "notebook"):
            self.notebook.select(self.tab_report)

        messagebox.showinfo(
            "Report dibuat",
            f"Summary:\n{self.summary_file}\n\nRejected:\n{self.rejected_file}"
        )


    # ============================================================
    # REPORT OPEN / EXPORT
    # ============================================================

    def find_latest_file(self, pattern):
        files = sorted(self.sessions_dir.glob(pattern), key=lambda p: p.stat().st_mtime if p.exists() else 0, reverse=True)
        return files[0] if files else None

    def open_last_summary(self):
        path = self.summary_file if self.summary_file and self.summary_file.exists() else self.find_latest_file("session_summary_*.csv")

        if not path:
            messagebox.showinfo("Belum ada summary", "Belum ada file session_summary.")
            return

        self.open_path(path)

    def open_last_rejected(self):
        path = self.rejected_file if self.rejected_file and self.rejected_file.exists() else self.find_latest_file("session_rejected_*.csv")

        if not path:
            messagebox.showinfo("Belum ada rejected log", "Belum ada file session_rejected.")
            return

        self.open_path(path)

    def read_csv_rows(self, path):
        if not path or not path.exists():
            return [], []

        with open(path, "r", newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        if not rows:
            return [], []

        return rows[0], rows[1:]

    def export_session_report_to_excel(self):
        summary_path = self.summary_file if self.summary_file and self.summary_file.exists() else self.find_latest_file("session_summary_*.csv")
        rejected_path = self.rejected_file if self.rejected_file and self.rejected_file.exists() else self.find_latest_file("session_rejected_*.csv")

        if not summary_path:
            messagebox.showwarning("Tidak ada summary", "Belum ada session_summary yang bisa diekspor.")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except Exception:
            messagebox.showerror(
                "openpyxl belum tersedia",
                "Export Excel membutuhkan openpyxl. Install dengan:\n\npython -m pip install openpyxl"
            )
            return

        excel_name = summary_path.stem.replace("session_summary_", "session_report_") + ".xlsx"
        excel_path = self.sessions_dir / excel_name

        wb = Workbook()

        # Sheet 1: Summary
        ws = wb.active
        ws.title = "Summary"
        header, rows = self.read_csv_rows(summary_path)
        ws.append(header)

        for row in rows:
            ws.append(row)

        # Sheet 2: Rejected
        ws2 = wb.create_sheet("Rejected")
        r_header, r_rows = self.read_csv_rows(rejected_path) if rejected_path else ([], [])
        if r_header:
            ws2.append(r_header)
            for row in r_rows:
                ws2.append(row)
        else:
            ws2.append(["Tidak ada rejected attempts"])

        # Sheet 3: Session Info
        ws3 = wb.create_sheet("Session Info")
        ws3.append(["Field", "Value"])
        ws3.append(["Session ID", self.session_id or ""])
        ws3.append(["Kelas", self.session_class or self.run_kelas.get()])
        ws3.append(["Mode", self.get_mode_name()])
        ws3.append(["Session Start", self.session_start.strftime("%Y-%m-%d %H:%M:%S") if self.session_start else ""])
        ws3.append(["Export Time", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
        ws3.append(["Summary CSV", str(summary_path)])
        ws3.append(["Rejected CSV", str(rejected_path) if rejected_path else ""])

        # Simple formatting
        fills = {
            "HADIR": PatternFill("solid", fgColor="D8F5D0"),
            "TIDAK HADIR": PatternFill("solid", fgColor="F8D7DA"),
        }

        for sheet in [ws, ws2, ws3]:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="E9ECEF")
            for column_cells in sheet.columns:
                length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
                sheet.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 45)

        status_col = 1
        for row in ws.iter_rows(min_row=2):
            status = row[status_col - 1].value
            if status in fills:
                for cell in row:
                    cell.fill = fills[status]

        wb.save(excel_path)
        self.last_excel_file = excel_path

        messagebox.showinfo("Export berhasil", f"Excel report berhasil dibuat:\n{excel_path}")
        self.open_path(excel_path)

    # ============================================================
    # FILE STATUS
    # ============================================================

    def required_files(self):
        kelas_lower = self.run_kelas.get().strip().lower()

        files = [
            self.project_root / "models" / f"face_recognition_mobilenet_{kelas_lower}.h5",
            self.project_root / "models" / f"class_indices_{kelas_lower}.json",
            self.backend_script,
        ]

        if self.antispoof_active_var.get():
            files.insert(0, self.project_root / "models" / "anti_spoofing_threshold.json")
            files.insert(0, self.project_root / "models" / "anti_spoofing_mobilenet_real_spoof.h5")

        return files

    def get_missing_files(self):
        return [str(p.relative_to(self.project_root)) for p in self.required_files() if not p.exists()]

    def check_required_files(self, show_popup=False):
        lines = []
        lines.append("PROJECT ROOT")
        lines.append(str(self.project_root))
        lines.append("")
        lines.append("REQUIRED FILES")
        lines.append("=" * 70)

        missing = []

        for path in self.required_files():
            status = "OK" if path.exists() else "MISSING"
            rel = path.relative_to(self.project_root)
            lines.append(f"[{status}] {rel}")

            if not path.exists():
                missing.append(str(rel))

        lines.append("")
        lines.append("TA MODEL NOTE")
        lines.append("=" * 70)
        lines.append("TA recognition model : models/face_recognition_mobilenet_ta.h5")
        lines.append("TA class indices     : models/class_indices_ta.json")
        lines.append("Backend              : src/final_attendance_lightweight_raspi.py")
        lines.append("Anti-spoofing        : mengikuti backend lama Raspberry Pi, tidak diubah")
        lines.append("")
        lines.append("DATABASE")
        lines.append("=" * 70)
        lines.append(f"Student CSV       : {self.student_csv}")
        lines.append(f"Raw attendance log: {self.attendance_log}")
        lines.append(f"Session folder    : {self.sessions_dir}")

        if hasattr(self, "status_text"):
            self.status_text.delete("1.0", "end")
            self.status_text.insert("end", "\n".join(lines))

        self.update_system_indicators()

        if show_popup:
            if missing:
                messagebox.showwarning("File belum lengkap", "File missing:\n\n" + "\n".join(missing))
            else:
                messagebox.showinfo("File lengkap", "Semua file penting untuk kelas yang dipilih sudah tersedia.")

    # ============================================================
    # OPEN HELPERS
    # ============================================================

    def open_path(self, path):
        try:
            if os.name == "nt":
                os.startfile(str(path))
            elif sys.platform == "darwin":
                subprocess.Popen(["open", str(path)])
            else:
                subprocess.Popen(["xdg-open", str(path)])
        except Exception as e:
            messagebox.showerror("Gagal membuka", str(e))

    def open_project_folder(self):
        self.open_path(self.project_root)

    def open_dataset_folder(self):
        self.dataset_dir.mkdir(parents=True, exist_ok=True)
        self.open_path(self.dataset_dir)

    def open_sessions_folder(self):
        self.sessions_dir.mkdir(parents=True, exist_ok=True)
        self.open_path(self.sessions_dir)

    def open_attendance_log(self):
        raw_log = self.get_current_raw_log()

        if not raw_log.exists():
            messagebox.showinfo("Belum ada log", f"File belum ada:\n{raw_log}")
            return

        self.open_path(raw_log)

    # ============================================================
    # CONSOLE
    # ============================================================

    def console_write(self, text):
        self.console.insert("end", text)
        self.console.see("end")

    def console_delete(self):
        self.console.delete("1.0", "end")


def main():
    app = AttendanceGUI()
    app.mainloop()


if __name__ == "__main__":
    main()
